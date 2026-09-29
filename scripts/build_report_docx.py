"""Build the editable, print-ready AgentFence report from REPORT.md.

Run with the Codex bundled Python environment containing python-docx.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "REPORT.md"
OUTPUT = ROOT / "docs" / "AgentFence_Project_Report.docx"
REPO = "https://github.com/codex-mohan/agentfence/blob/main/"
INK = RGBColor(20, 20, 20)
MUTED = RGBColor(95, 95, 95)
LIGHT = "F4F4F4"
BORDER = "D9D9D9"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    item = tc_pr.find(qn("w:tcBorders"))
    if item is None:
        item = OxmlElement("w:tcBorders")
        tc_pr.append(item)
    for side in ("top", "left", "bottom", "right"):
        edge = item.find(qn(f"w:{side}"))
        if edge is None:
            edge = OxmlElement(f"w:{side}")
            item.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), BORDER)


def add_hyperlink(paragraph, label: str, href: str) -> None:
    label = label.strip("`")
    if not href.startswith(("http://", "https://")):
        href = REPO + href.removeprefix("./")
    rel = paragraph.part.relate_to(
        href,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "333333")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


TOKEN = re.compile(r"(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")


def inline(paragraph, text: str) -> None:
    for part in TOKEN.split(text):
        if not part:
            continue
        link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", part)
        if link:
            add_hyperlink(paragraph, link.group(1), link.group(2))
        elif part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif part.startswith("*") and part.endswith("*"):
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:tblHeader")
    flag.set(qn("w:val"), "true")
    tr_pr.append(flag)


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:cantSplit")
    tr_pr.append(flag)


def add_table(doc, source_lines: list[str]) -> None:
    parsed = [[cell.strip() for cell in line.strip().strip("|").split("|")]
              for line in source_lines]
    rows = [parsed[0]] + [row for row in parsed[2:] if row]
    column_count = len(parsed[0])
    table = doc.add_table(rows=len(rows), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    available = 16.2
    if column_count == 2:
        proportions = [0.38, 0.62]
    elif column_count == 3:
        proportions = [0.22, 0.39, 0.39]
    else:
        proportions = [1 / column_count] * column_count
    for row_idx, values in enumerate(rows):
        keep_row_together(table.rows[row_idx])
        for col_idx in range(column_count):
            cell = table.cell(row_idx, col_idx)
            cell.width = Cm(available * proportions[col_idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            borders(cell)
            if row_idx == 0:
                shade(cell, "282828")
            elif row_idx % 2 == 0:
                shade(cell, "F5F5F5")
            paragraph = cell.paragraphs[0]
            paragraph.style = "Table Text"
            inline(paragraph, values[col_idx] if col_idx < len(values) else "")
            if row_idx == 0:
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
    set_repeat_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_flow(doc) -> None:
    p = doc.add_paragraph(style="Diagram Label")
    p.add_run("Figure 1  AgentFence decision flow")
    flows = [
        ("Before a tool runs", "Agent host  →  Adapter or Guard  →  Engine and trusted policy  →  Allow, deny, or review"),
        ("After a tool returns", "Covered executor  →  Result inspection and optional Laya score  →  Release or withhold"),
        ("Evidence", "Each decision and observed outcome  →  SQLite event log  →  Live local dashboard"),
    ]
    table = doc.add_table(rows=3, cols=2)
    table.autofit = False
    for i, (label, content) in enumerate(flows):
        keep_row_together(table.rows[i])
        left, right = table.rows[i].cells
        left.width, right.width = Cm(4.2), Cm(12)
        for cell in (left, right):
            borders(cell)
            shade(cell, "F2F2F2" if i % 2 == 0 else "FFFFFF")
        left.paragraphs[0].style = "Table Text"
        left.paragraphs[0].add_run(label).bold = True
        right.paragraphs[0].style = "Table Text"
        right.paragraphs[0].add_run(content)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code(doc, lines: list[str]) -> None:
    for line in lines:
        p = doc.add_paragraph(style="Code Block")
        p.paragraph_format.keep_with_next = True
        p.add_run(line if line else " ")
    if lines:
        doc.paragraphs[-1].paragraph_format.keep_with_next = False
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def style_document(doc) -> None:
    section = doc.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21)
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    section.header_distance = Cm(0.95)
    section.footer_distance = Cm(0.95)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.16

    title = doc.styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(13)
    title_ppr = title.element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size, before, after in (
        ("Heading 1", 14, 16, 7),
        ("Heading 2", 11.5, 11, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name, size, color, space in (
        ("Table Text", 8.6, INK, 2),
        ("Code Block", 8.5, INK, 0),
        ("Diagram Label", 9, MUTED, 5),
    ):
        style = doc.styles.add_style(style_name, 1)
        style.font.name = "Consolas" if style_name == "Code Block" else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_after = Pt(space)
        style.paragraph_format.line_spacing = 1.08

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("AGENTFENCE  /  CAPSTONE PROJECT REPORT")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("AgentFence   •   ")
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def build() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    style_document(doc)
    title = doc.add_paragraph(style="Title")
    title.add_run("AgentFence Capstone Project Report")
    title_ppr = title._p.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    subtitle = doc.add_paragraph()
    subtitle.add_run("Architecture, live demonstration, evaluation, and limits").italic = True
    subtitle.paragraph_format.space_after = Pt(14)

    index = 1
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith('<p align="center"'):
            icon = doc.add_paragraph()
            icon.alignment = WD_ALIGN_PARAGRAPH.CENTER
            icon.add_run().add_picture(str(ROOT / "assets" / "agentfence-icon.png"), width=Inches(0.87))
            index += 1
            continue
        if line.startswith("## "):
            heading = line[3:]
            if heading.startswith("1. "):
                doc.add_page_break()
            heading = heading.replace(":", "").replace(",", "")
            doc.add_paragraph(heading, style="Heading 1")
            index += 1
            continue
        if line.startswith("### "):
            doc.add_paragraph(line[4:].replace(":", ""), style="Heading 2")
            index += 1
            continue
        if line.startswith("| "):
            block = []
            while index < len(lines) and lines[index].startswith("|"):
                block.append(lines[index])
                index += 1
            add_table(doc, block)
            continue
        if line.startswith("```"):
            language = line[3:]
            block = []
            index += 1
            while index < len(lines) and not lines[index].startswith("```"):
                block.append(lines[index])
                index += 1
            if language == "mermaid":
                add_flow(doc)
            else:
                add_code(doc, block)
            index += 1
            continue
        numbered = re.match(r"^\d+\. (.*)$", line)
        if line.startswith("- ") or numbered:
            if line.startswith("- "):
                p = doc.add_paragraph(style="List Bullet")
                inline(p, line[2:])
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.65)
                p.paragraph_format.first_line_indent = Cm(-0.55)
                p.add_run(line.split(".", 1)[0] + ".  ")
                inline(p, numbered.group(1))
            index += 1
            continue
        block = [line]
        index += 1
        while index < len(lines) and lines[index].strip() and not re.match(
            r"^(#{1,3} |\| |```|- |\d+\. |<p)", lines[index]
        ):
            block.append(lines[index].strip())
            index += 1
        p = doc.add_paragraph()
        inline(p, " ".join(block))

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
